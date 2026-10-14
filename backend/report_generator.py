
from pathlib import Path
from datetime import datetime
from xml.sax.saxutils import escape
import re

from docx import Document
from docx.shared import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable
)


def clean_filename(topic):
    name = re.sub(r'[<>:"/\\|?*]', "_", topic)
    name = re.sub(r"\s+", "_", name).strip("._ ")
    return name[:100] or "Research_Report"


def clean_topic(topic):
    return topic.rstrip("\\/ \t\r\n")


def markdown_to_plain_text(text):
    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    text = re.sub(
        r"\[([^\]]+)\]\((https?://[^)]+)\)",
        r"\1 (\2)",
        text
    )
    text = re.sub(r"[*_`~]", "", text)
    return text.strip()


def remove_duplicate_references(report):
    lines = report.splitlines()
    result = []
    seen_urls = set()
    in_references = False

    for line in lines:
        if re.match(r"^\s{0,3}#{1,3}\s+(references|sources)\s*$",
                    line, re.IGNORECASE):
            in_references = True
            result.append(line)
            continue

        if in_references:
            urls = re.findall(r"https?://[^\s)\]>]+", line)
            if urls:
                normalized = urls[0].rstrip(".,;")
                if normalized in seen_urls:
                    continue
                seen_urls.add(normalized)

        result.append(line)

    return "\n".join(result)


def generate_docx(topic, report, output_path):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    title = document.add_heading(clean_topic(topic), level=0)
    title.alignment = 1

    paragraph = document.add_paragraph()
    paragraph.add_run("Research Report | Generated ").bold = True
    paragraph.add_run(datetime.now().strftime("%d %B %Y"))

    document.add_paragraph("")

    for line in report.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith("### "):
            document.add_heading(line[4:], level=3)

        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)

        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)

        elif re.match(r"^[-*+]\s+", line):
            text = re.sub(r"^[-*+]\s+", "", line)
            document.add_paragraph(
                markdown_to_plain_text(text),
                style="List Bullet"
            )

        elif re.match(r"^\d+[.)]\s+", line):
            text = re.sub(r"^\d+[.)]\s+", "", line)
            document.add_paragraph(
                markdown_to_plain_text(text),
                style="List Number"
            )

        else:
            document.add_paragraph(markdown_to_plain_text(line))

    document.save(str(output_path))


def generate_pdf(topic, report, output_path):
    styles = getSampleStyleSheet()

    styles.add(ParagraphStyle(
        name="ReportTitleCustom",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#17365D"),
        spaceAfter=18
    ))

    styles.add(ParagraphStyle(
        name="ReportHeading1Custom",
        parent=styles["Heading1"],
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#17365D"),
        spaceBefore=14,
        spaceAfter=7,
        keepWithNext=True
    ))

    styles.add(ParagraphStyle(
        name="ReportHeading2Custom",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#285A84"),
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True
    ))

    styles.add(ParagraphStyle(
        name="ReportBodyCustom",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        spaceAfter=7,
        splitLongWords=1
    ))

    story = [
        Paragraph("Research Report", styles["ReportTitleCustom"]),
        Paragraph(
            "<b>Topic:</b> " + escape(clean_topic(topic)),
            styles["ReportBodyCustom"]
        ),
        Paragraph(
            "<b>Generated:</b> "
            + datetime.now().strftime("%d %B %Y"),
            styles["ReportBodyCustom"]
        ),
        Spacer(1, 8),
        HRFlowable(
            width="100%",
            color=colors.HexColor("#D5DDE5")
        ),
        Spacer(1, 12)
    ]

    for line in report.splitlines():
        line = line.strip()

        if not line:
            story.append(Spacer(1, 3))
            continue

        if line.startswith("### "):
            story.append(Paragraph(
                escape(line[4:]),
                styles["ReportHeading2Custom"]
            ))

        elif line.startswith("## "):
            story.append(Paragraph(
                escape(line[3:]),
                styles["ReportHeading1Custom"]
            ))

        elif line.startswith("# "):
            story.append(Paragraph(
                escape(line[2:]),
                styles["ReportHeading1Custom"]
            ))

        elif re.match(r"^[-*+]\s+", line):
            item = re.sub(r"^[-*+]\s+", "", line)
            story.append(Paragraph(
                "&#8226; " + escape(markdown_to_plain_text(item)),
                styles["ReportBodyCustom"]
            ))

        elif re.match(r"^\d+[.)]\s+", line):
            item = re.sub(r"^\d+[.)]\s+", "", line)
            story.append(Paragraph(
                escape(markdown_to_plain_text(item)),
                styles["ReportBodyCustom"]
            ))

        else:
            story.append(Paragraph(
                escape(markdown_to_plain_text(line)),
                styles["ReportBodyCustom"]
            ))

    def add_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawCentredString(
            A4[0] / 2,
            1.1 * cm,
            f"Page {doc.page}"
        )
        canvas.restoreState()

    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=f"Research Report - {clean_topic(topic)}",
        author="Multi-Agent AI Research System"
    )

    pdf.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number
    )


def generate_all_formats(topic, report, reports_folder):
    reports_folder = Path(reports_folder)
    reports_folder.mkdir(parents=True, exist_ok=True)

    topic = clean_topic(topic)
    report = remove_duplicate_references(report)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = f"{clean_filename(topic)}_{timestamp}"

    md_path = reports_folder / f"{base}.md"
    pdf_path = reports_folder / f"{base}.pdf"
    docx_path = reports_folder / f"{base}.docx"

    md_path.write_text(report, encoding="utf-8")
    generate_pdf(topic, report, pdf_path)
    generate_docx(topic, report, docx_path)

    return md_path, pdf_path, docx_path
